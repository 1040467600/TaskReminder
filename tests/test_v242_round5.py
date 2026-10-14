"""第五轮新增测试：app_config 边界、tasks_page 排序/列配置、repository 状态流转、excel_io 行过滤、db 事务提交。"""
from __future__ import annotations

import json
from datetime import datetime, timedelta

import pytest

from task_reminder import app_config, db, repository as repo
from task_reminder.models import TaskStatus, fmt, html_to_plain


# ======================================================================
# app_config.py — None / JSON 字符串边界
# ======================================================================
class TestAppConfigEdgeCases:
    def test_set_none_stores_null_get_returns_none(self):
        app_config.set("none_key", None)
        assert app_config.get("none_key") is None

    def test_set_json_looking_string_gets_parsed(self):
        """字符串 '123' 存储原样，但 get 优先 json.loads → 返回 int 123。"""
        app_config.set("str_num", "123")
        assert app_config.get("str_num") == 123

    def test_set_non_json_string_returns_as_is(self):
        app_config.set("plain_str", "hello world")
        assert app_config.get("plain_str") == "hello world"

    def test_set_empty_string(self):
        app_config.set("empty_str", "")
        assert app_config.get("empty_str") == ""

    def test_set_bool_true_false(self):
        app_config.set("flag", False)
        assert app_config.get("flag") is False
        app_config.set("flag", True)
        assert app_config.get("flag") is True

    def test_first_run_done_default_false(self):
        assert app_config.get("first_run_done") is False

    def test_overwrite_existing_key(self):
        app_config.set("overwrite_me", 1)
        app_config.set("overwrite_me", "two")
        assert app_config.get("overwrite_me") == "two"


# ======================================================================
# tasks_page.py — 排序文本 & 非排序列 & 列配置持久化
# ======================================================================
class TestTasksPageSortAndColumns:
    def _mk_task(self):
        now = datetime.now()
        return repo.add_task("排序测试", "张三", "",
                             fmt(now + timedelta(days=1)), fmt(now + timedelta(hours=1)))

    def test_current_sort_text_default(self, qtbot):
        from task_reminder.ui.tasks_page import TasksPage
        self._mk_task()
        page = TasksPage()
        qtbot.addWidget(page)
        assert "默认" in page.current_sort_text()

    def test_current_sort_text_asc(self, qtbot):
        from task_reminder.ui.tasks_page import TasksPage
        self._mk_task()
        page = TasksPage()
        qtbot.addWidget(page)
        dl = page.model.column_ids.index("deadline")
        page._on_header_clicked(dl)
        txt = page.current_sort_text()
        assert "截止时间" in txt
        assert "升序" in txt
        assert "▲" in txt

    def test_current_sort_text_desc(self, qtbot):
        from task_reminder.ui.tasks_page import TasksPage
        self._mk_task()
        page = TasksPage()
        qtbot.addWidget(page)
        dl = page.model.column_ids.index("deadline")
        page._on_header_clicked(dl)  # asc
        page._on_header_clicked(dl)  # desc
        txt = page.current_sort_text()
        assert "降序" in txt
        assert "▼" in txt

    def test_non_sortable_column_no_change(self, qtbot):
        """actions 列不可排序：点击不改变 sort_key/sort_dir。"""
        from task_reminder.ui.tasks_page import TasksPage
        self._mk_task()
        page = TasksPage()
        qtbot.addWidget(page)
        actions_idx = page.model.column_ids.index("actions")
        page._on_header_clicked(actions_idx)
        assert page.sort_dir == ""
        assert page.sort_key == "created_at"

    def test_toggle_column_persists(self, qtbot):
        from task_reminder.ui.tasks_page import TasksPage
        self._mk_task()
        page = TasksPage()
        qtbot.addWidget(page)
        name_idx = page.model.column_ids.index("name")
        page._toggle_column(name_idx, False)
        # 持久化到 app_config
        state = app_config.get("column_state")
        assert state is not None
        found = [c for c in state if c["id"] == "name"]
        assert found and found[0]["visible"] is False

    def test_reset_columns_restores_defaults(self, qtbot):
        from task_reminder.ui.tasks_page import TasksPage
        self._mk_task()
        page = TasksPage()
        qtbot.addWidget(page)
        notes_idx = page.model.column_ids.index("notes")
        # 先显示备注列
        page._toggle_column(notes_idx, True)
        assert not page.view.horizontalHeader().isSectionHidden(notes_idx)
        # 重置后备注列恢复隐藏
        page._reset_columns()
        assert page.view.horizontalHeader().isSectionHidden(notes_idx)


# ======================================================================
# repository.py — update_task 状态流转 & content_plain 同步
# ======================================================================
class TestRepositoryUpdateTaskFlow:
    def test_update_syncs_content_plain(self):
        now = datetime.now()
        tid = repo.add_task("内容同步", "张三", "<b>粗体</b>",
                            fmt(now + timedelta(days=1)), fmt(now + timedelta(hours=1)))
        t = repo.get_task(tid)
        assert t.content_plain == "粗体"
        # 更新内容，content_plain 应同步
        repo.update_task(tid, name="内容同步", assignee="张三",
                         content="<i>斜体</i>", deadline=fmt(now + timedelta(days=1)),
                         reminder_time=fmt(now + timedelta(hours=1)),
                         status=TaskStatus.NOT_STARTED, notes="")
        assert repo.get_task(tid).content_plain == "斜体"

    def test_update_to_done_sets_triggered(self):
        now = datetime.now()
        tid = repo.add_task("完成测试", "张三", "",
                            fmt(now + timedelta(days=1)), fmt(now + timedelta(hours=1)))
        assert repo.get_task(tid).triggered == 0
        repo.update_task(tid, name="完成测试", assignee="张三", content="",
                         deadline=fmt(now + timedelta(days=1)),
                         reminder_time=fmt(now + timedelta(hours=1)),
                         status=TaskStatus.DONE, notes="")
        assert repo.get_task(tid).triggered == 1

    def test_update_from_done_to_not_started_future_reminder_resets_triggered(self):
        now = datetime.now()
        future_rm = fmt(now + timedelta(hours=2))
        tid = repo.add_task("回退测试", "张三", "",
                            fmt(now + timedelta(days=1)), future_rm)
        # 先标记完成 → triggered=1
        repo.update_task(tid, name="回退测试", assignee="张三", content="",
                         deadline=fmt(now + timedelta(days=1)),
                         reminder_time=future_rm,
                         status=TaskStatus.DONE, notes="")
        assert repo.get_task(tid).triggered == 1
        # 从完成回退到未开始，提醒在未来 → triggered 重置为 0
        repo.update_task(tid, name="回退测试", assignee="张三", content="",
                         deadline=fmt(now + timedelta(days=1)),
                         reminder_time=future_rm,
                         status=TaskStatus.NOT_STARTED, notes="")
        assert repo.get_task(tid).triggered == 0

    def test_update_from_done_to_not_started_past_reminder_keeps_triggered(self):
        """提醒时间已过，从完成回退到未开始时 triggered 保持 1（不再触发提醒）。"""
        now = datetime.now()
        past_rm = fmt(now - timedelta(hours=1))
        tid = repo.add_task("回退过提醒", "张三", "",
                            fmt(now + timedelta(days=1)), past_rm)
        repo.update_task(tid, name="回退过提醒", assignee="张三", content="",
                         deadline=fmt(now + timedelta(days=1)),
                         reminder_time=past_rm,
                         status=TaskStatus.DONE, notes="")
        assert repo.get_task(tid).triggered == 1
        repo.update_task(tid, name="回退过提醒", assignee="张三", content="",
                         deadline=fmt(now + timedelta(days=1)),
                         reminder_time=past_rm,
                         status=TaskStatus.NOT_STARTED, notes="")
        assert repo.get_task(tid).triggered == 1


# ======================================================================
# excel_io.py — read_rows 行过滤
# ======================================================================
class TestExcelReadRowsFiltering:
    def test_skips_meta_row(self, tmp_path):
        from openpyxl import Workbook
        from task_reminder.excel_io import HEADERS, read_rows
        path = str(tmp_path / "meta.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(["任务A", "张三", "内容", "2026-09-10 10:00", "2026-09-09 09:00",
                   "未开始", "", "", "其他"])
        ws.append([f"导出时间：{datetime.now()}    共 1 条"])
        wb.save(path)
        rows = read_rows(path)
        assert len(rows) == 1
        assert rows[0]["name"] == "任务A"

    def test_skips_completely_empty_row(self, tmp_path):
        from openpyxl import Workbook
        from task_reminder.excel_io import HEADERS, read_rows
        path = str(tmp_path / "empty.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(["任务A", "张三", "内容", "2026-09-10 10:00", "2026-09-09 09:00",
                   "未开始", "", "", "其他"])
        ws.append(["", "", "", "", "", "", "", "", ""])  # 全空行
        ws.append(["任务B", "李四", "内容2", "2026-09-11 10:00", "2026-09-10 09:00",
                   "未开始", "", "", "其他"])
        wb.save(path)
        rows = read_rows(path)
        assert len(rows) == 2
        assert [r["name"] for r in rows] == ["任务A", "任务B"]

    def test_skips_sparse_row_without_name_assignee_deadline(self, tmp_path):
        """名称、执行人、截止时间全空的稀疏行跳过；有名称则保留。"""
        from openpyxl import Workbook
        from task_reminder.excel_io import HEADERS, read_rows
        path = str(tmp_path / "sparse.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(["任务A", "张三", "内容", "2026-09-10 10:00", "2026-09-09 09:00",
                   "未开始", "", "", "其他"])
        ws.append(["", "", "", "", "", "", "", "", ""])  # 全空行跳过
        ws.append(["孤立名称", "", "", "", "", "", "", "", ""])  # 有名称 → 保留
        wb.save(path)
        rows = read_rows(path)
        names = [r["name"] for r in rows]
        assert "任务A" in names
        assert "孤立名称" in names
        # 全空行被跳过
        assert all(r["name"] != "" for r in rows)


# ======================================================================
# db.py — transaction 提交
# ======================================================================
class TestDbTransactionCommit:
    def test_transaction_commits_on_success(self):
        # 用一个临时表验证事务提交后数据持久
        c = db.get_conn()
        c.execute("CREATE TABLE IF NOT EXISTS tx_test (v INTEGER)")
        c.execute("DELETE FROM tx_test")
        with db.transaction() as conn:
            conn.execute("INSERT INTO tx_test(v) VALUES (?)", (42,))
        # 事务外可见
        row = c.execute("SELECT v FROM tx_test").fetchone()
        assert row[0] == 42
        c.execute("DROP TABLE tx_test")
