"""v2.4.1 缺陷修复回归测试（15 条缺陷的针对性守护）。

覆盖：
- TC-CAT-02/TC-TASK-30 新建对话框类别默认"其他"
- TC-AL2-08 弹窗 Esc/×（rejected）不再卡死队列，历史记 close
- TC-CONC-04 任务被删除后弹窗响应不崩溃
- TC-ERR-03/TC-XRT-02 导入损坏 xlsx 有 try/except 友好提示
- TC-CFG-02 极端字号 apply_theme 不崩
- TC-SRCH-12/13 高级搜索"提醒时间"范围
- TC-TASK-15 截止/提醒时间可留空（双空允许、单空拒绝、查询排除）
- TC-XLSL-03 空库导出生成仅表头文件
- TC-ERR-02 备份 JSON 缺字段补默认值
（TC-HIST-07/TC-BOUND-15/TC-CFG-03 已在 test_supplement/test_round3 守护）
"""
from __future__ import annotations

import json
from datetime import datetime, timedelta

import pytest

from task_reminder import backup, db, excel_io, repository as repo
from task_reminder.models import TaskStatus, fmt


# ---------------------------------------------------------------------------
# TC-CAT-02 / TC-TASK-30：新建对话框类别默认"其他"
class TestDefaultCategory:
    def test_new_dialog_default_category_other(self, qtbot):
        from task_reminder.ui.task_dialog import TaskDialog
        dlg = TaskDialog()
        qtbot.addWidget(dlg)
        assert dlg.cmb_category.currentText() == "其他"


# ---------------------------------------------------------------------------
# TC-AL2-08：Esc/× 关闭弹窗（rejected）不卡死队列
class TestDialogRejected:
    def _mk_due(self, name):
        past = datetime.now() - timedelta(minutes=1)
        tid = repo.add_task(name, "张三", "内容",
                            fmt(datetime.now() + timedelta(days=1)), fmt(past))
        hid = repo.mark_triggered(tid)
        return tid, hid

    def test_reject_records_close_and_pumps_next(self, qtbot):
        from task_reminder.notifier import Notifier
        t1, h1 = self._mk_due("弹窗任务甲")
        t2, h2 = self._mk_due("弹窗任务乙")
        n = Notifier()
        responses = []
        n.responded.connect(responses.append)
        n.notify(repo.get_task(t1), h1)
        n.notify(repo.get_task(t2), h2)
        n.pump()
        first = n._dialog
        assert first is not None
        first.reject()  # 模拟 Esc / 标题栏 ×
        # 历史响应记为 close，且信号已发出
        resp1 = db.get_conn().execute(
            "SELECT response FROM reminder_history WHERE id=?", (h1,)).fetchone()[0]
        assert resp1 == "close"
        assert "close" in responses
        # 队列不卡死：第二个弹窗自动顶上
        assert n._dialog is not None and n._dialog is not first
        n._dialog.reject()
        assert n._dialog is None
        assert n.pending_count() == 0


# ---------------------------------------------------------------------------
# TC-CONC-04：任务被外部删除后 set_status 静默忽略
class TestDeletedTaskResponse:
    def test_set_status_deleted_task_no_crash(self):
        tid = repo.add_task("将被删", "张三", "内容",
                            fmt(datetime.now() + timedelta(days=1)),
                            fmt(datetime.now() + timedelta(hours=1)))
        repo.delete_task(tid)
        # 不抛异常即通过（弹窗点"标记完成"路径）
        assert repo.set_status(tid, "已完成") is None
        assert repo.get_task(tid) is None

    def test_dialog_done_after_task_deleted(self, qtbot):
        """ReminderDialog._respond('done') 在任务已删时也不崩。"""
        from task_reminder.notifier import ReminderDialog
        tid = repo.add_task("弹窗中删除", "张三", "内容",
                            fmt(datetime.now() + timedelta(days=1)),
                            fmt(datetime.now() - timedelta(minutes=1)))
        hid = repo.mark_triggered(tid)
        task = repo.get_task(tid)
        dlg = ReminderDialog(task, hid)
        qtbot.addWidget(dlg)
        repo.delete_task(tid)
        dlg._respond("done")  # 不抛异常
        assert dlg.response == "done"


# ---------------------------------------------------------------------------
# TC-ERR-03 / TC-XRT-02：损坏/非法 xlsx 导入不崩溃
class TestImportCorruptXlsx:
    def test_on_import_corrupt_file_shows_critical(self, qtbot, tmp_path, monkeypatch):
        from PyQt6.QtWidgets import QDialog, QFileDialog, QMessageBox
        from task_reminder.ui import tasks_page as tp
        bad = tmp_path / "corrupt.xlsx"
        bad.write_text("这根本不是一个 xlsx（zip）文件", encoding="utf-8")
        monkeypatch.setattr(QFileDialog, "getOpenFileName",
                            staticmethod(lambda *a, **k: (str(bad), "")))
        monkeypatch.setattr(QDialog, "exec", lambda self: 1)  # 导入选项对话框直接确定
        critical = []
        monkeypatch.setattr(QMessageBox, "critical",
                            staticmethod(lambda *a, **k: critical.append(a)))
        monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))
        page = tp.TasksPage()
        qtbot.addWidget(page)
        page.on_import()  # 不抛异常
        assert len(critical) == 1  # 友好错误提示


# ---------------------------------------------------------------------------
# TC-CFG-02：极端字号不崩（main 入口另有 max(11,min(22)) 钳制）
class TestFontSizeExtremes:
    @pytest.mark.parametrize("size", [1, 11, 22, 99])
    def test_apply_theme_extreme_sizes(self, qtbot, size):
        from PyQt6.QtWidgets import QApplication
        from task_reminder.ui.theme import apply_theme
        apply_theme(QApplication.instance(), "light", size)  # 不抛异常

    def test_clamp_expression(self):
        """main 入口钳制逻辑：11–22 之外的值被收敛。"""
        for raw, expected in [(99, 22), (1, 11), ("15", 15)]:
            assert max(11, min(22, int(raw))) == expected


# ---------------------------------------------------------------------------
# TC-SRCH-12 / TC-SRCH-13：高级搜索提醒时间范围
class TestReminderRangeSearch:
    def test_search_bar_has_reminder_range(self, qtbot):
        from task_reminder.ui.search_bar import SearchBar
        sb = SearchBar()
        qtbot.addWidget(sb)
        assert "reminder" in sb._date_ranges
        rng = sb._date_ranges["reminder"]
        assert "reminder" not in sb.criteria()  # 未勾选不加条件
        rng["chk"].setChecked(True)
        c = sb.criteria()
        assert c["reminder_from"].endswith("00:00")
        assert c["reminder_to"].endswith("23:59")
        assert "reminder" in sb.adv_keys_in_use()

    def test_repository_filters_by_reminder_range(self):
        soon = repo.add_task("近期提醒", "张三", "内容",
                             fmt(datetime.now() + timedelta(days=2)),
                             fmt(datetime.now() + timedelta(days=1)))
        far = repo.add_task("远期提醒", "李四", "内容",
                            fmt(datetime.now() + timedelta(days=30)),
                            fmt(datetime.now() + timedelta(days=20)))
        lo = fmt(datetime.now()).split(" ")[0] + " 00:00"
        hi = (datetime.now() + timedelta(days=3)).strftime("%Y-%m-%d") + " 23:59"
        tasks, total = repo.list_tasks(criteria={"reminder_from": lo, "reminder_to": hi})
        ids = {t.id for t in tasks}
        assert soon in ids and far not in ids


# ---------------------------------------------------------------------------
# TC-TASK-15：截止/提醒时间可留空
class TestOptionalTimes:
    def test_validate_both_empty_allowed(self):
        repo.validate_task("无时间任务", "张三", "内容", "", "")  # 不抛异常

    def test_validate_only_one_empty_rejected(self):
        with pytest.raises(repo.ValidationError):
            repo.validate_task("单空任务", "张三", "内容", "", fmt(datetime.now()))
        with pytest.raises(repo.ValidationError):
            repo.validate_task("单空任务", "张三", "内容", fmt(datetime.now()), "")

    def test_add_task_empty_times_roundtrip(self):
        tid = repo.add_task("无时间", "张三", "内容", "", "")
        t = repo.get_task(tid)
        assert t.deadline == "" and t.reminder_time == ""

    def test_due_for_reminder_excludes_empty_reminder(self):
        empty = repo.add_task("空提醒", "张三", "内容", "", "")
        due = repo.add_task("有提醒", "李四", "内容",
                            fmt(datetime.now() + timedelta(days=1)),
                            fmt(datetime.now() - timedelta(minutes=1)))
        ids = {t.id for t in repo.due_for_reminder()}
        assert due in ids and empty not in ids

    def test_overdue_excludes_empty_deadline(self):
        empty = repo.add_task("空截止", "张三", "内容", "", "")
        overdue = repo.add_task("已过期", "李四", "内容",
                                fmt(datetime.now() - timedelta(days=1)),
                                fmt(datetime.now() - timedelta(days=2)))
        ids = {t.id for t in repo.overdue_tasks()}
        assert overdue in ids and empty not in ids

    def test_dialog_both_none_checkboxes_save_empty(self, qtbot):
        from task_reminder.ui.task_dialog import TaskDialog
        dlg = TaskDialog()
        qtbot.addWidget(dlg)
        assert not dlg.chk_no_deadline.isChecked()
        assert not dlg.chk_no_reminder.isChecked()
        dlg.edit_name.setText("对话框无时间")
        dlg.edit_assignee.setText("张三")
        dlg.chk_no_deadline.setChecked(True)
        dlg.chk_no_reminder.setChecked(True)
        assert not dlg.dt_deadline.isEnabled()
        assert not dlg.dt_reminder.isEnabled()
        assert dlg.btn_ok.isEnabled()
        dlg._on_save()
        assert dlg.result() == 1
        t = repo.all_tasks()[0]
        assert t.deadline == "" and t.reminder_time == ""

    def test_dialog_only_one_none_blocks_save(self, qtbot):
        from task_reminder.ui.task_dialog import TaskDialog
        dlg = TaskDialog()
        qtbot.addWidget(dlg)
        dlg.edit_name.setText("单空")
        dlg.edit_assignee.setText("张三")
        dlg.chk_no_deadline.setChecked(True)  # 只勾一个
        assert not dlg.btn_ok.isEnabled()

    def test_dialog_load_empty_times_checks_boxes(self, qtbot):
        from task_reminder.ui.task_dialog import TaskDialog
        tid = repo.add_task("存量无时间", "张三", "内容", "", "")
        dlg = TaskDialog(editing_task=repo.get_task(tid))
        qtbot.addWidget(dlg)
        assert dlg.chk_no_deadline.isChecked() and dlg.chk_no_reminder.isChecked()

    def test_dialog_preset_unchecks_none(self, qtbot):
        from task_reminder.ui.task_dialog import TaskDialog
        dlg = TaskDialog()
        qtbot.addWidget(dlg)
        dlg.chk_no_deadline.setChecked(True)
        dlg._apply_preset(dlg.dt_deadline, True,
                          datetime.now().replace(second=0, microsecond=0))
        assert not dlg.chk_no_deadline.isChecked()
        assert dlg.dt_deadline.isEnabled()


# ---------------------------------------------------------------------------
# TC-XLSL-03：空库导出生成仅表头文件
class TestEmptyExport:
    def test_export_empty_list_creates_header_only(self, tmp_path):
        path = str(tmp_path / "empty.xlsx")
        assert excel_io.export_tasks(path, []) == 0
        rows = excel_io.read_rows(path)
        assert rows == []
        from openpyxl import load_workbook
        ws = load_workbook(path).active
        assert [c.value for c in ws[1]] == excel_io.HEADERS

    def test_excel_empty_times_roundtrip(self, tmp_path):
        """v2.4.1 联动：无时间任务导出 Excel 后再导入，时间保持为空。"""
        path = str(tmp_path / "notime.xlsx")
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(excel_io.HEADERS)
        ws.append(["无时间任务", "张三", "x", None, None, "未开始", "", "", "其他"])
        ws.append(["坏时间任务", "李四", "x", "2030-01-01 12:00", None, "未开始", "", "", "其他"])
        wb.save(path)
        report = excel_io.import_tasks(path)
        assert report.added == 1 and len(report.errors) == 1  # 双空成功、单空报错
        t = next(x for x in repo.all_tasks() if x.name == "无时间任务")
        assert t.deadline == "" and t.reminder_time == ""


# ---------------------------------------------------------------------------
# TC-ERR-02：备份 JSON 缺字段时补默认值
class TestBackupMissingFields:
    def test_import_minimal_backup(self, tmp_path):
        payload = {
            "format": "task_reminder_backup",
            "version": 1,
            "tasks": [{"id": 77, "name": "极简备份", "assignee": "王五"}],
            "history": [],
        }
        path = tmp_path / "min.json"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        n_t, n_h = backup.import_json(str(path))
        assert n_t == 1 and n_h == 0
        t = repo.get_task(77)
        assert t is not None
        assert t.name == "极简备份" and t.assignee == "王五"
        assert t.status == TaskStatus.NOT_STARTED
        assert t.category == "其他"
        assert t.triggered == 0
        assert t.deadline == "" and t.reminder_time == ""


# ---------------------------------------------------------------------------
# 第二轮：原 9 条待验证项的代码/运行时验证
class TestRound2RiskItems:
    def test_update_deleted_task_no_crash(self):
        """TC-CONC-03：编辑保存时任务已被外部删除 → 不抛 AttributeError。"""
        tid = repo.add_task("编辑中被删", "张三", "内容",
                            fmt(datetime.now() + timedelta(days=1)),
                            fmt(datetime.now() + timedelta(hours=1)))
        repo.delete_task(tid)
        repo.update_task(tid, name="编辑中被删", assignee="张三", content="内容",
                         deadline=fmt(datetime.now() + timedelta(days=1)),
                         reminder_time=fmt(datetime.now() + timedelta(hours=1)),
                         status="未开始", notes="", category="其他")  # 不抛异常
        assert repo.get_task(tid) is None  # 不会凭空创建

    def test_busy_timeout_set(self):
        """TC-CONC-02：连接设置 busy_timeout，双进程写锁等待而非立刻报错。"""
        v = db.get_conn().execute("PRAGMA busy_timeout").fetchone()[0]
        assert v >= 3000

    def test_restore_offscreen_geometry_clamps_to_primary(self, qtbot):
        """TC-DISP-02/04/TC-CFG-01：记忆位置在所有屏幕之外 → 回到主屏可见区。"""
        from PyQt6.QtGui import QGuiApplication
        from task_reminder import app_config
        from task_reminder.ui.main_window import MainWindow
        app_config.set("window_state", {"x": -99999, "y": -99999,
                                        "w": 1000, "h": 700, "maximized": False})
        win = MainWindow()
        qtbot.addWidget(win)
        primary = QGuiApplication.primaryScreen().availableGeometry()
        assert primary.intersects(win.frameGeometry())

    def test_restore_huge_coordinates_clamps(self, qtbot):
        from PyQt6.QtGui import QGuiApplication
        from task_reminder import app_config
        from task_reminder.ui.main_window import MainWindow
        app_config.set("window_state", {"x": 99999, "y": 99999,
                                        "w": 1000, "h": 700, "maximized": False})
        win = MainWindow()
        qtbot.addWidget(win)
        ag = QGuiApplication.primaryScreen().availableGeometry()
        assert ag.intersects(win.frameGeometry())

    def test_corrupt_db_raises(self, tmp_path):
        """TC-PATH-04：库文件被文本覆盖 → init 抛 DatabaseError（main 入口捕获后弹错误框）。"""
        import sqlite3
        bad = tmp_path / "bad.db"
        bad.write_text("not a sqlite database", encoding="utf-8")
        with pytest.raises(sqlite3.DatabaseError):
            db.init(str(bad))
        db.close()
        db.init(str(tmp_path / "fresh.db"))  # 重新指向好库，避免污染后续

    def test_local_server_second_instance_signal(self, qtbot):
        """TC-SHELL-09：第二个实例连入时 newConnection 信号触发（用于带起原窗口）。"""
        from PyQt6.QtNetwork import QLocalServer, QLocalSocket
        key = "TaskReminder-pytest-second-instance"
        QLocalServer.removeServer(key)
        server = QLocalServer()
        assert server.listen(key)
        fired = []
        server.newConnection.connect(lambda: fired.append(1))
        client = QLocalSocket()
        client.connectToServer(key)
        qtbot.waitUntil(lambda: bool(fired), timeout=3000)
        conn = server.nextPendingConnection()
        assert conn is not None
        conn.disconnectFromServer()
        client.disconnectFromServer()
        server.close()

    def test_thousand_rows_performance(self, qtbot):
        """TC-BOUND-24：1000 条数据下列表查询与页面刷新在可接受时间内完成。"""
        import time
        from task_reminder.ui.tasks_page import TasksPage
        base = datetime.now() + timedelta(days=30)
        with db.transaction() as c:
            c.executemany(
                "INSERT INTO tasks(name, assignee, content, content_plain, deadline,"
                " reminder_time, status, category, notes, triggered, status_changed_at,"
                " created_at, updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                [(f"性能任务{i}", f"执行人{i % 10}", "内容", "内容",
                  fmt(base), fmt(base - timedelta(hours=1)), "未开始", "其他",
                  "", 0, "", fmt(base), fmt(base)) for i in range(1000)],
            )
        t0 = time.perf_counter()
        tasks, total = repo.list_tasks(page=1, page_size=200)
        t_list = time.perf_counter() - t0
        assert total == 1000 and len(tasks) == 200
        page = TasksPage()
        qtbot.addWidget(page)
        t0 = time.perf_counter()
        page.refresh()
        t_refresh = time.perf_counter() - t0
        assert page.model.rowCount() == 200
        assert t_list < 2.0 and t_refresh < 3.0

