"""Excel 导入导出单测。"""
from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from task_reminder import excel_io, repository as repo
from task_reminder.models import fmt


def make_xlsx(path, rows, headers=excel_io.HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def sample_task(**kw):
    now = datetime.now()
    defaults = dict(name="导出测试任务", assignee="李四", content="<p>正文<b>加粗</b></p>",
                    deadline=fmt(now), reminder=fmt(datetime.now().replace(year=now.year - 1)),
                    status="进行中", notes="备注内容")
    defaults.update(kw)
    tid = repo.add_task(defaults["name"], defaults["assignee"], defaults["content"],
                        defaults["deadline"], defaults["reminder"], defaults["status"],
                        defaults["notes"])
    return repo.get_task(tid)


class TestExport:
    def test_export_and_headers(self, tmp_path):
        sample_task()
        path = str(tmp_path / "out.xlsx")
        n = excel_io.export_tasks(path, repo.all_tasks())
        assert n == 1
        wb = load_workbook(path)
        ws = wb.active
        assert [c.value for c in ws[1]] == excel_io.HEADERS
        row2 = [c.value for c in ws[2]]
        assert row2[0] == "导出测试任务"
        assert "<b>" not in row2[2]          # 内容已转纯文本
        assert "加粗" in row2[2]

    def test_export_unicode_path(self, tmp_path):
        sample_task()
        path = str(tmp_path / "任务导出 中文.xlsx")
        assert excel_io.export_tasks(path, repo.all_tasks()) == 1


class TestImport:
    def test_roundtrip(self, tmp_path):
        sample_task()
        src = str(tmp_path / "a.xlsx")
        excel_io.export_tasks(src, repo.all_tasks())

        # 清库后导入
        repo.delete_tasks([t.id for t in repo.all_tasks()])
        report = excel_io.import_tasks(src)
        assert report.total == 1 and report.added == 1 and not report.errors
        t = repo.all_tasks()[0]
        assert t.name == "导出测试任务" and t.assignee == "李四"
        assert t.status == "进行中" and t.notes == "备注内容"

    def test_import_validations(self, tmp_path):
        now = datetime.now()
        rows = [
            ["", "张三", "x", fmt(now), "2020-01-01 10:00", "未开始", "", ""],   # 名称缺失（下限 1 字符）
            ["合法的任务名称", "张三", "x", "", "2020-01-01 10:00", "未开始", "", ""],  # 缺截止
            ["合法的任务名称", "张三", "x", "bad-time", "2020-01-01 10:00", "未开始", "", ""],  # 时间错
            ["合法的任务名称", "张三", "x", fmt(now), "2030-01-01 10:00", "未开始", "", ""],  # 提醒晚于截止
            ["合法的任务名称", "张三", "x", fmt(now), "2020-01-01 10:00", "未知状态", "", ""],  # 状态回退默认
            ["合法的任务名称2", "张三", "x", fmt(now), "2020-01-01 10:00", "未开始", "", ""],  # OK
        ]
        path = str(tmp_path / "in.xlsx")
        make_xlsx(path, rows)
        report = excel_io.import_tasks(path)
        assert report.added == 2                 # 未知状态回退默认后可导入
        assert report.total == 6 and len(report.errors) == 4
        assert report.errors[0].startswith("第 2 行")

    def test_conflict_skip_and_overwrite(self, tmp_path):
        t = sample_task(name="冲突测试任务")
        now = datetime.now()
        rows = [[t.name, t.assignee, "覆盖后的内容", t.deadline, "2020-01-01 10:00", "未开始", "", ""]]
        path = str(tmp_path / "dup.xlsx")
        make_xlsx(path, rows)

        r1 = excel_io.import_tasks(path, strategy="skip")
        assert r1.skipped == 1 and r1.updated == 0
        assert repo.get_task(t.id).plain_text() != "覆盖后的内容"

        r2 = excel_io.import_tasks(path, strategy="overwrite")
        assert r2.updated == 1 and r2.skipped == 0
        assert repo.get_task(t.id).plain_text() == "覆盖后的内容"

    def test_excel_datetime_cells(self, tmp_path):
        dt = datetime(2026, 9, 15, 8, 30)
        path = str(tmp_path / "dt.xlsx")
        make_xlsx(path, [["日期单元格任务", "王五", "x", dt, datetime(2026, 9, 14, 8, 30),
                          "未开始", "", ""]])
        report = excel_io.import_tasks(path)
        assert report.added == 1
        t = repo.all_tasks()[0]
        assert t.deadline == "2026-09-15 08:30"

    def test_meta_row_ignored(self, tmp_path):
        now = datetime.now()
        path = str(tmp_path / "meta.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(excel_io.HEADERS)
        ws.append(["元信息任务名称", "赵六", "x", fmt(now), "2020-01-01 10:00", "未开始", "", ""])
        ws.append([])
        ws.append(["导出时间：2026-09-05 10:00    共 1 条    由 任务提醒助手 导出"])
        wb.save(path)
        report = excel_io.import_tasks(path)
        assert report.total == 1 and report.added == 1
